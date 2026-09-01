import tensorflow, joblib, matplotlib.pyplot
from sklearn.svm import LinearSVC
from openpyxl import load_workbook
from sklearn.metrics import accuracy_score, confusion_matrix, ConfusionMatrixDisplay

#Load model
model = joblib.load(r"C:\Users\Asnie\Downloads\MMW-HAT\MMW-HAT\MMW-HAT-Release\example_4_offline_proc\SVM_Experiment_Augmented.joblib")

#Test data
test_folder  = r"C:\Users\Asnie\Downloads\MMW-HAT\MMW-HAT\MMW-HAT-Release\example_4_offline_proc\Features - Experiment + Augmented - SVM\Test"

features_test = []
activity_test = []

test_files = [file for file in tensorflow.io.gfile.listdir(test_folder) if file.endswith(".xlsx")]
for file in test_files:
    file_path = test_folder + "\\" + file
    workbook = load_workbook(file_path, data_only = True)
    sheet = workbook[workbook.sheetnames[-1]]
    features = []
    for row in sheet.iter_rows(min_row = 1, values_only = True):
        features.append(row[1])
    features_test.append(features)
    activity_test.append(sheet.title)


#Prediction
predicted_features = model.predict(features_test)

#Accuracy
accuracy = accuracy_score(activity_test, predicted_features)
print("SVM test accuracy:", accuracy)

#Actual vs predicted activity
for i, (actual, predicted) in enumerate(zip(activity_test, predicted_features)):
    print(f"Actual activity = {actual} | Predicted activity = {predicted}")

#Confusion matrix
class_names = ["Cyclist", "Pedestrian"]
confusion_matrix_svm = confusion_matrix(activity_test, predicted_features)
confusion_matrix_svm_plot = ConfusionMatrixDisplay(confusion_matrix = confusion_matrix_svm, display_labels = class_names)
confusion_matrix_svm_plot.plot(cmap = "RdBu", text_kw = {"color": "white", "fontsize": 18})
matplotlib.pyplot.title("SVM Confusion Matrix", fontsize = 20)
confusion_matrix_svm_plot.ax_.tick_params(axis = "both", labelsize = 15)
confusion_matrix_svm_plot.ax_.xaxis.label.set_size(15)
confusion_matrix_svm_plot.ax_.yaxis.label.set_size(15)
matplotlib.pyplot.savefig(r"C:\Users\Asnie\Downloads\MMW-HAT\MMW-HAT\MMW-HAT-Release\example_4_offline_proc\SVM Confusion Matrix - Experiment + Augmented.png", bbox_inches = "tight")
matplotlib.pyplot.show()

#Importance of features
importance = abs(model.coef_).mean(axis=0)
importance_pct = 100 * importance / importance.sum()
feature_names = [row[0] for row in sheet.iter_rows(min_row=1, values_only=True)]
for name, pct in zip(feature_names, importance_pct):
    print(f"{name}: {pct:.2f}%")
